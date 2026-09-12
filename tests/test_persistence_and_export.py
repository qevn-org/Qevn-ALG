"""Unit tests for SQLite Lead Repository and Export Service (CSV, JSON, 5-sheet XLSX)."""

import json

import openpyxl
import pytest

from linkedin_intelligence.db.repository import SqliteLeadRepository
from linkedin_intelligence.models.lead import (
    Company,
    ContactInfo,
    Lead,
    Person,
    SearchRunRecord,
)
from linkedin_intelligence.services.export_service import (
    export_leads_csv,
    export_leads_json,
    export_leads_xlsx,
)


@pytest.fixture
def test_repo(tmp_path):
    db_file = tmp_path / "test_leads.db"
    return SqliteLeadRepository(db_path=db_file)


def test_repository_upsert_and_filtering(test_repo):
    lead1 = Lead(
        id="lead-1",
        score=94.0,
        temperature="HOT",
        status="new",
        person=Person(name="Amit Patel", job_title="CTO", is_decision_maker=True),
        company=Company(name="Apex Systems", industry="Healthcare & MedTech", location="Bangalore, India"),
        contacts=[ContactInfo(type="business_email", value="amit@apex.com", confidence=0.95, verification_status="verified_public")],
        primary_email="amit@apex.com",
    )
    lead2 = Lead(
        id="lead-2",
        score=58.0,
        temperature="COOL",
        status="new",
        person=Person(name="Sneha Rao", job_title="Developer", is_decision_maker=False),
        company=Company(name="Beta Corp", industry="Fintech & Financial Services", location="Mumbai, India"),
    )

    test_repo.upsert_leads([lead1, lead2])

    # 1. Fetch single
    fetched = test_repo.get_lead("lead-1")
    assert fetched is not None
    assert fetched.company.name == "Apex Systems"
    assert fetched.score == 94.0

    # 2. Filter by temperature
    hot_leads = test_repo.list_leads(temperature="HOT")
    assert len(hot_leads) == 1
    assert hot_leads[0].id == "lead-1"

    # 3. Filter by industry
    health_leads = test_repo.list_leads(industry="Healthcare")
    assert len(health_leads) == 1
    assert health_leads[0].company.name == "Apex Systems"

    # 4. Filter by email
    email_leads = test_repo.list_leads(has_email=True)
    assert len(email_leads) == 1
    assert email_leads[0].id == "lead-1"

    # 5. Search term
    searched = test_repo.list_leads(search_term="Beta")
    assert len(searched) == 1
    assert searched[0].company.name == "Beta Corp"


def test_repository_outcome_recording(test_repo):
    lead = Lead(
        id="lead-3",
        score=88.0,
        company=Company(name="Gamma Labs"),
        feature_vector={"hiring_intent": 85.0, "urgency": 90.0},
    )
    test_repo.upsert_lead(lead)

    # Record outcome action
    outcome = test_repo.record_outcome(
        lead_id="lead-3",
        action="won",
        label=1,
        features=lead.feature_vector,
        notes="Closed $50k engagement",
    )
    assert outcome.action == "won"
    assert outcome.label == 1

    # Check outcomes dataset
    dataset = test_repo.get_outcomes_dataset()
    assert len(dataset) == 1
    assert dataset[0]["action"] == "won"
    assert dataset[0]["label"] == 1
    assert dataset[0]["features"]["hiring_intent"] == 85.0

    # Updated lead status in repo
    updated_lead = test_repo.get_lead("lead-3")
    assert updated_lead.outcome == "won"
    assert updated_lead.status == "won"


def test_search_run_persistence(test_repo):
    run = SearchRunRecord(
        run_id="run-test-01",
        user_query="healthcare technology companies in India",
        start_time="2026-09-12T10:00:00Z",
        result_count=50,
        qualified_count=12,
    )
    test_repo.save_search_run(run)

    runs = test_repo.get_search_runs()
    assert len(runs) == 1
    assert runs[0].run_id == "run-test-01"
    assert runs[0].result_count == 50


def test_export_csv_and_json():
    leads = [
        Lead(
            id="exp-1",
            score=91.0,
            temperature="HOT",
            person=Person(name="Vikram Seth", job_title="Founder"),
            company=Company(name="Delta Health", industry="Healthcare"),
        )
    ]

    csv_bytes = export_leads_csv(leads)
    assert len(csv_bytes) > 0
    csv_text = csv_bytes.decode("utf-8")
    assert "Delta Health" in csv_text
    assert "Vikram Seth" in csv_text

    json_str = export_leads_json(leads)
    assert "Delta Health" in json_str
    parsed = json.loads(json_str)
    assert isinstance(parsed, list)
    assert parsed[0]["id"] == "exp-1"


def test_export_multi_sheet_xlsx():
    leads = [
        Lead(
            id="exp-2",
            score=89.5,
            temperature="WARM",
            person=Person(name="Meera Kapoor", job_title="VP Product"),
            company=Company(name="Zeta AI", industry="AI / SaaS"),
            contacts=[ContactInfo(type="business_email", value="meera@zeta.ai", confidence=0.96, verification_status="verified_public")],
            evidence=["post-zeta-1"],
            why_detected=["Hiring lead engineers"],
            feature_vector={"hiring_intent": 90.0, "urgency": 80.0},
        )
    ]

    xlsx_bytes = export_leads_xlsx(leads)
    assert len(xlsx_bytes) > 0

    import io
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = wb.sheetnames
    assert sheet_names == ["Leads", "Contacts", "Signals", "Evidence", "Scoring"]

    # Verify sheet contents
    ws_leads = wb["Leads"]
    assert ws_leads.max_row >= 2
    assert ws_leads["A1"].value is not None

    ws_contacts = wb["Contacts"]
    assert ws_contacts.max_row >= 2
    assert "meera@zeta.ai" in [cell.value for row in ws_contacts.iter_rows() for cell in row]

    ws_scoring = wb["Scoring"]
    assert ws_scoring.max_row >= 2
